"""Gemini-nin birbaşa qəbul etmədiyi ofis fayllarını (Word, Excel, CSV) strukturlu mətnə çevirir.

PDF və şəkillər modelə olduğu kimi göndərilir. Burada şablon yoxdur: cədvəllər sətir-sətir saxlanılır,
məna və sütunları model özü müəyyən edir.
"""
import csv
import io
import zipfile
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree

MAX_TEXT_CHARS = 2_000_000
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

class ConversionError(ValueError):
    pass

def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).replace('\t', ' ').replace('\n', ' ').strip()

def rows_to_text(rows):
    lines = []
    for row in rows:
        cells = [cell_text(value) for value in row]
        while cells and not cells[-1]:
            cells.pop()
        if cells:
            lines.append('\t'.join(cells))
    return '\n'.join(lines)

def xlsx_to_text(raw):
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ConversionError('Excel faylı açılmadı; fayl zədələnib və ya şifrəlidir.') from None
    try:
        parts = [f'=== Page {index}: sheet "{sheet.title}" ===\n{rows_to_text(sheet.iter_rows(values_only=True))}'
                 for index, sheet in enumerate(workbook.worksheets, start=1)]
    finally:
        workbook.close()
    return '\n\n'.join(parts)

def xls_to_text(raw):
    import xlrd
    try:
        workbook = xlrd.open_workbook(file_contents=raw)
    except Exception:
        raise ConversionError('Köhnə Excel (.xls) faylı açılmadı; faylı .xlsx və ya PDF kimi saxlayın.') from None
    parts = []
    for index, sheet in enumerate(workbook.sheets(), start=1):
        rows = ([sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows))
        parts.append(f'=== Page {index}: sheet "{sheet.name}" ===\n{rows_to_text(rows)}')
    return '\n\n'.join(parts)

def csv_to_text(raw):
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        raise ConversionError('CSV UTF-8 formatında olmalıdır.') from None
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel
    return '=== Page 1 ===\n' + rows_to_text(csv.reader(io.StringIO(text), dialect))

def docx_to_text(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            root = ElementTree.fromstring(archive.read('word/document.xml'))
    except (zipfile.BadZipFile, KeyError, ElementTree.ParseError):
        raise ConversionError('Word faylı açılmadı; fayl zədələnib və ya şifrəlidir.') from None

    def paragraph(node):
        return ''.join(text.text or '' for text in node.iter(W + 't')).strip()

    body = root.find(W + 'body')
    if body is None:
        raise ConversionError('Word faylında mətn tapılmadı.')
    blocks = []
    for node in body:
        if node.tag == W + 'p':
            text = paragraph(node)
            if text:
                blocks.append(text)
        elif node.tag == W + 'tbl':
            rows = [[' '.join(filter(None, (paragraph(p) for p in cell.iter(W + 'p')))) for cell in row.findall(W + 'tc')]
                    for row in node.iter(W + 'tr')]
            blocks.append('[table]\n' + rows_to_text(rows) + '\n[/table]')
    return '=== Page 1 ===\n' + '\n'.join(blocks)

CONVERTERS = {'.xlsx': xlsx_to_text, '.xls': xls_to_text, '.csv': csv_to_text, '.docx': docx_to_text}

def is_office(name):
    return Path(name).suffix.lower() in CONVERTERS

def to_text(raw, name):
    text = CONVERTERS[Path(name).suffix.lower()](raw)
    content = [line for line in text.splitlines() if line.strip() and not line.startswith('=== Page ')]
    if not content or all(line in ('[table]', '[/table]') for line in content):
        raise ConversionError('Faylda oxunacaq məzmun tapılmadı.')
    if len(text) > MAX_TEXT_CHARS:
        raise ConversionError('Faylın məzmunu çox böyükdür; sənədi hissələrə bölün.')
    return f'[Converted from {Path(name).suffix.lower()} file "{Path(name).name}". Tab separates cells.]\n{text}'
